#!/usr/bin/env python3
"""
Seed service_mappings and payment_type_mappings tables
from COA_Quickbooks_matched.xlsx

This script can create the database if it doesn't exist, or update an existing one.
"""

import sqlite3
import openpyxl
import uuid
import json
import sys
import argparse
from datetime import datetime
from pathlib import Path

# Paths
SCRIPT_DIR = Path(__file__).parent
DEFAULT_DB_PATH = SCRIPT_DIR / '../src/electron-app/accounting.db'
EXCEL_PATH = SCRIPT_DIR / '../data/raw/COA_Quickbooks_matched.xlsx'

# Payment type mapping configuration
PAYMENT_MAPPINGS = {
    'Alle Rewards': {
        'category': 'vendor_receivable',
        'clearing_account': '1200 · Accounts Receivable:1220 · Vendor Receivables:Allē Rewards'
    },
    'Aspire Awards': {
        'category': 'vendor_receivable',
        'clearing_account': '1200 · Accounts Receivable:1220 · Vendor Receivables:Aspire Awards'
    },
    'Cherry': {
        'category': 'financing',
        'clearing_account': '1200 · Accounts Receivable:1220 · Vendor Receivables:Cherry Financing'
    },
    'Amex': {
        'category': 'credit_card',
        'clearing_account': '1030 · Merchant Clearing'
    },
    'Visa': {
        'category': 'credit_card',
        'clearing_account': '1030 · Merchant Clearing'
    },
    'MasterCard': {
        'category': 'credit_card',
        'clearing_account': '1030 · Merchant Clearing'
    },
    'Discover': {
        'category': 'credit_card',
        'clearing_account': '1030 · Merchant Clearing'
    },
    'Cash': {
        'category': 'cash',
        'clearing_account': '1010 · Cash - Operating'
    },
    'Check': {
        'category': 'check',
        'clearing_account': '1020 · Undeposited Funds'
    },
    'Square Gift Card': {
        'category': 'gift_card',
        'clearing_account': '2300 · Gift Card Liability'
    },
    'Reward Points': {
        'category': 'loyalty',
        'clearing_account': '2310 · Customer Loyalty Points'
    },
    'Client Bank': {
        'category': 'loyalty',
        'clearing_account': '2320 · Client Bank Liability'
    }
}

def read_excel_data():
    """Read service and payment mappings from Excel file"""
    print('📊 Reading Excel file...')

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # Read service mappings
    services = []
    if 'emr_service_items' in wb.sheetnames:
        ws = wb['emr_service_items']
        headers = [cell.value for cell in ws[1]]

        for row in ws.iter_rows(min_row=2, values_only=True):
            services.append(dict(zip(headers, row)))

    # Read payment mappings
    payments = []
    if 'emr_payment_types' in wb.sheetnames:
        ws = wb['emr_payment_types']
        headers = [cell.value for cell in ws[1]]

        for row in ws.iter_rows(min_row=2, values_only=True):
            payments.append(dict(zip(headers, row)))

    print(f'  ✓ Found {len(services)} service mappings')
    print(f'  ✓ Found {len(payments)} payment types\n')

    return services, payments

def seed_service_mappings(conn, services):
    """Insert service mappings into database"""
    print('📥 Inserting service mappings...')

    cursor = conn.cursor()

    # Clear existing data (optional)
    cursor.execute('DELETE FROM service_mappings')

    insert_query = '''
        INSERT INTO service_mappings (
            id, emr_service_name, qb_item_name, qb_item_hierarchy,
            asset_account, income_account, tax_code, is_active
        ) VALUES (?, ?, ?, ?, ?, ?, ?, 1)
    '''

    count = 0
    for svc in services:
        # Skip if service name is null
        service_name = svc.get('Service/Product')
        if not service_name:
            continue

        item_name = svc.get('Matched_Item') or ''

        # Extract hierarchy (parent category) from item name
        item_hierarchy = item_name.split(':')[0] if ':' in item_name else ''

        cursor.execute(insert_query, (
            str(uuid.uuid4()),
            service_name,
            item_name,
            item_hierarchy,
            svc.get('Asset Account'),
            svc.get('Account') or '',
            svc.get('Tax Code') or 'Non'
        ))
        count += 1

    print(f'  ✓ Inserted {count} service mappings\n')
    return count

def seed_payment_mappings(conn, payments):
    """Insert payment type mappings into database"""
    print('📥 Inserting payment type mappings...')

    cursor = conn.cursor()

    # Clear existing data (optional)
    cursor.execute('DELETE FROM payment_type_mappings')

    insert_query = '''
        INSERT INTO payment_type_mappings (
            id, payment_type, category, clearing_account, is_active
        ) VALUES (?, ?, ?, ?, 1)
    '''

    count = 0
    for pmt in payments:
        payment_type = pmt.get('Payment Type')
        if not payment_type:
            continue

        # Get mapping configuration or use default
        mapping = PAYMENT_MAPPINGS.get(payment_type, {
            'category': 'other',
            'clearing_account': '1030 · Merchant Clearing'
        })

        cursor.execute(insert_query, (
            str(uuid.uuid4()),
            payment_type,
            mapping['category'],
            mapping['clearing_account']
        ))
        count += 1

    print(f'  ✓ Inserted {count} payment type mappings\n')
    return count

def add_audit_log(conn, service_count, payment_count):
    """Add audit log entry for seeding operation"""
    cursor = conn.cursor()

    details = json.dumps({
        'service_count': service_count,
        'payment_count': payment_count,
        'source_file': 'COA_Quickbooks_matched.xlsx',
        'timestamp': datetime.now().isoformat()
    })

    cursor.execute('''
        INSERT INTO audit_log (action, entity_type, details)
        VALUES (?, ?, ?)
    ''', ('seed_database', 'mappings', details))

def create_database_schema(conn):
    """Create database tables if they don't exist"""
    print('📋 Creating database schema...')

    schema = '''
    CREATE TABLE IF NOT EXISTS service_mappings (
      id                 TEXT PRIMARY KEY,
      emr_service_name   TEXT NOT NULL UNIQUE,
      qb_item_name       TEXT NOT NULL,
      qb_item_hierarchy  TEXT,
      asset_account      TEXT,
      income_account     TEXT NOT NULL,
      tax_code           TEXT,
      is_active          INTEGER DEFAULT 1,
      created_at         TEXT DEFAULT (datetime('now')),
      updated_at         TEXT DEFAULT (datetime('now'))
    );

    CREATE INDEX IF NOT EXISTS idx_service_mappings_emr_service ON service_mappings(emr_service_name);
    CREATE INDEX IF NOT EXISTS idx_service_mappings_active ON service_mappings(is_active);

    CREATE TABLE IF NOT EXISTS payment_type_mappings (
      id                 TEXT PRIMARY KEY,
      payment_type       TEXT UNIQUE NOT NULL,
      category           TEXT NOT NULL,
      clearing_account   TEXT NOT NULL,
      is_active          INTEGER DEFAULT 1,
      created_at         TEXT DEFAULT (datetime('now'))
    );

    CREATE INDEX IF NOT EXISTS idx_payment_type ON payment_type_mappings(payment_type);

    CREATE TABLE IF NOT EXISTS audit_log (
      id                 INTEGER PRIMARY KEY AUTOINCREMENT,
      timestamp          TEXT DEFAULT (datetime('now')),
      action             TEXT NOT NULL,
      entity_type        TEXT,
      entity_id          TEXT,
      details            TEXT,
      user_id            TEXT DEFAULT 'primary_user'
    );

    CREATE INDEX IF NOT EXISTS idx_audit_log_timestamp ON audit_log(timestamp);
    CREATE INDEX IF NOT EXISTS idx_audit_log_action ON audit_log(action);
    '''

    conn.executescript(schema)
    print('  ✓ Schema created\n')

def main():
    parser = argparse.ArgumentParser(description='Seed database with mappings from Excel')
    parser.add_argument('--db', type=str, help=f'Database path (default: {DEFAULT_DB_PATH})')
    parser.add_argument('--create', action='store_true', help='Create database if it doesn\'t exist')
    args = parser.parse_args()

    db_path = Path(args.db) if args.db else DEFAULT_DB_PATH

    print('🌱 Starting database seeding...\n')

    # Check if database exists
    if not db_path.exists():
        if args.create:
            print(f'📝 Creating new database at: {db_path}')
            db_path.parent.mkdir(parents=True, exist_ok=True)
        else:
            print(f'❌ Database not found at: {db_path}')
            print('   Use --create flag to create a new database, or specify path with --db')
            print('   Or start the Electron app first to create the database.')
            return 1

    # Check if Excel file exists
    if not EXCEL_PATH.exists():
        print(f'❌ Excel file not found at: {EXCEL_PATH}')
        return 1

    # Read Excel data
    services, payments = read_excel_data()

    # Connect to database
    print(f'💾 Connecting to database: {db_path}')
    conn = sqlite3.connect(str(db_path))
    print('  ✓ Connected\n')

    # Create schema if needed
    create_database_schema(conn)

    try:
        # Seed service mappings
        service_count = seed_service_mappings(conn, services)

        # Seed payment mappings
        payment_count = seed_payment_mappings(conn, payments)

        # Add audit log
        add_audit_log(conn, service_count, payment_count)

        # Commit changes
        print('💾 Committing changes...')
        conn.commit()
        print('  ✓ Committed\n')

        print('✅ Seeding complete!')
        print(f'   📊 {service_count} service mappings')
        print(f'   💳 {payment_count} payment type mappings')

        return 0

    except Exception as e:
        print(f'❌ Error during seeding: {e}')
        conn.rollback()
        return 1

    finally:
        conn.close()

if __name__ == '__main__':
    exit(main())
